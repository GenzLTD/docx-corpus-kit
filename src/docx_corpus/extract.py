"""Text extraction from .docx, with a robust fallback ladder.

- docx_text_plain : python-docx, fast & reliable (paragraphs + tables). Used for
  bulk indexing where speed matters more than rich formatting.
- docx_text_md    : markitdown, high-quality markdown. Runs in a thread with a
  timeout so a single pathological file cannot hang a batch job.
- docx_text_best  : prefers markdown; on timeout/error/oversized file falls back
  to plain text. This is what callers should use.
"""
import os
import threading


def docx_text_plain(path):
    """Fast, dependency-light plain-text extraction (python-docx)."""
    try:
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    except Exception as e:  # pragma: no cover - defensive
        return "[extract-error: %s]" % e


def docx_text_md(path, timeout=60):
    """High-quality markdown via markitdown. Returns None on timeout/error."""
    try:
        from markitdown import MarkItDown
        box = {}

        def _run():
            try:
                box["v"] = MarkItDown().convert(path).text_content or ""
            except Exception as e:  # pragma: no cover - defensive
                box["e"] = e

        th = threading.Thread(target=_run, daemon=True)
        th.start()
        th.join(timeout)
        if th.is_alive() or "e" in box:
            return None
        return box.get("v", "")
    except Exception:  # pragma: no cover - defensive
        return None


def docx_text_best(path, timeout=60, big_threshold=800 * 1024):
    """Best-effort extraction: markdown when possible, plain text otherwise.

    Files larger than ``big_threshold`` skip markitdown entirely (it can hang on
    very large documents) and go straight to the fast plain-text path.
    """
    try:
        if os.path.getsize(path) <= big_threshold:
            md = docx_text_md(path, timeout)
            if md is not None:
                return md
    except OSError:  # pragma: no cover - defensive
        pass
    base = os.path.splitext(os.path.basename(path))[0]
    return "# %s\n\n%s" % (base, docx_text_plain(path))
